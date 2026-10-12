# -*- coding: utf-8 -*-
"""확정된 findings JSON 을 풋팅결과 엑셀로 만든다.

사용법
    python build_report.py 03_findings.json -o "풋팅결과_보고서명.xlsx"

입력 스키마
{
 "meta": {"title": "...", "subtitle": "...", "note": "..."},
 "summary": {"S": "...", "A": "...", "B": "...", "C": "...", "V": "...",
             "conclusion": "여러 줄 결론"},
 "findings": [{"sev":"S|A|B|C", "no":1, "type":"결론·근거",
               "loc":"슬라이드4 (주1)", "issue":"...", "basis":"...", "fix":"..."}],
 "tieout":  [{"no":1, "item":"...", "report":29997417222, "source":"...",
              "source_value":29997417222, "diff":0, "basis":"...", "flag":false}],
 "verified":[{"no":1, "item":"...", "detail":"...", "source":"...", "result":"일치"}]
}

summary 의 S/A/B/C/V 는 각 등급 한 줄 요약. 건수는 findings 에서 자동으로 센다.
"""
import argparse
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

F = "맑은 고딕"
NAVY = "002060"
HDR = PatternFill("solid", fgColor="DCE6F1")
FILL = {"S": PatternFill("solid", fgColor="FFC7CE"),
        "A": PatternFill("solid", fgColor="FFEB9C"),
        "B": PatternFill("solid", fgColor="DDEBF7"),
        "C": PatternFill("solid", fgColor="EDEDED"),
        "V": PatternFill("solid", fgColor="E2EFDA")}
_thin = Side(style="thin", color="BFBFBF")
BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
SEV_ORDER = {"S": 0, "A": 1, "B": 2, "C": 3}
SEV_NAME = {"S": "발행 전 필수 수정 — 결론·근거",
            "A": "내용 보강·확인 필요",
            "B": "표 형식·레이아웃",
            "C": "문구·오타·띄어쓰기"}


def put(ws, r, vals, size=9):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font = Font(name=F, size=size)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BD


def header(ws, r, cols, widths):
    put(ws, r, cols)
    for i in range(1, len(cols) + 1):
        c = ws.cell(r, i)
        c.font = Font(name=F, size=10, bold=True, color=NAVY)
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def est_height(texts, widths, floor=32, cap=200):
    lines = 1
    for t, w in zip(texts, widths):
        t = str(t or "")
        lines = max(lines, len(t) / max(w * 0.92, 8) + t.count("\n"))
    return max(floor, min(lines * 14 + 8, cap))


def sheet_summary(wb, data):
    ws = wb.active
    ws.title = "요약"
    ws.sheet_view.showGridLines = False
    m = data.get("meta", {})
    ws["A1"] = m.get("title", "풋팅 결과")
    ws["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    for i, key in enumerate(("subtitle", "note"), start=2):
        if m.get(key):
            ws.cell(i, 1, m[key]).font = Font(name=F, size=9)

    counts = {k: 0 for k in "SABC"}
    for f in data.get("findings", []):
        counts[f["sev"]] = counts.get(f["sev"], 0) + 1
    counts["V"] = len(data.get("verified", []))

    s = data.get("summary", {})
    r = 5
    header(ws, r, ["중요도", "구분", "건수", "핵심 내용"], [8, 26, 7, 112])
    for i, k in enumerate(["S", "A", "B", "C", "V"], 1):
        name = SEV_NAME.get(k, "검증 완료 (이상 없음)")
        put(ws, r + i, [k if k != "V" else "✓", name, counts.get(k, 0), s.get(k, "")])
        ws.cell(r + i, 1).fill = FILL[k]
        ws.cell(r + i, 1).font = Font(name=F, size=10, bold=True)
        ws.cell(r + i, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r + i, 3).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r + i].height = est_height([s.get(k, "")], [112], 40, 90)

    r2 = r + 7
    ws.cell(r2, 1, "핵심 결론").font = Font(name=F, size=11, bold=True, color=NAVY)
    ws.cell(r2 + 1, 1, s.get("conclusion", "")).font = Font(name=F, size=9)
    ws.cell(r2 + 1, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=r2 + 1, start_column=1, end_row=r2 + 1, end_column=4)
    ws.row_dimensions[r2 + 1].height = est_height([s.get("conclusion", "")], [112], 60, 320)


def sheet_findings(wb, data):
    ws = wb.create_sheet("지적사항")
    ws.sheet_view.showGridLines = False
    ws["A1"] = ("지적사항 — 중요도 순 "
                "(S: 발행 전 필수 / A: 내용 보강·확인 / B: 표 형식·레이아웃 / C: 문구·오타)")
    ws["A1"].font = Font(name=F, size=12, bold=True, color=NAVY)
    widths = [8, 5, 11, 24, 58, 62, 60]
    header(ws, 3, ["중요도", "No", "유형", "위치", "지적 내용", "근거 · 계산", "수정안"], widths)

    items = sorted(data.get("findings", []),
                   key=lambda x: (SEV_ORDER.get(x["sev"], 9), x.get("no", 0)))
    r = 4
    for i, f in enumerate(items, 1):
        row = [f["sev"], f.get("no", i), f.get("type", ""), f.get("loc", ""),
               f.get("issue", ""), f.get("basis", ""), f.get("fix", "")]
        put(ws, r, row)
        ws.cell(r, 1).fill = FILL[f["sev"]]
        ws.cell(r, 1).font = Font(name=F, size=10, bold=True)
        for c in (1, 2, 3):
            ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center",
                                                wrap_text=True)
        ws.row_dimensions[r].height = est_height(row[4:], widths[4:])
        r += 1
    ws.freeze_panes = "A4"
    if r > 4:
        ws.auto_filter.ref = "A3:G%d" % (r - 1)


def sheet_tieout(wb, data):
    rows = data.get("tieout", [])
    if not rows:
        return
    ws = wb.create_sheet("숫자대사")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "숫자 대사 명세 — 보고서 기재액 ↔ 조서·원장"
    ws["A1"].font = Font(name=F, size=12, bold=True, color=NAVY)
    widths = [5, 30, 18, 46, 18, 14, 58]
    header(ws, 3, ["No", "항목", "보고서 기재", "대사 원천", "원천 금액", "차이", "대사 근거"], widths)
    r = 4
    for i, x in enumerate(rows, 1):
        put(ws, r, [x.get("no", i), x.get("item", ""), x.get("report"),
                    x.get("source", ""), x.get("source_value"), x.get("diff"),
                    x.get("basis", "")])
        for c in (3, 5, 6):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ws.cell(r, c).number_format = "#,##0" if abs(v) > 100 else "0.000000"
                ws.cell(r, c).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        if x.get("flag"):
            for c in range(1, 8):
                ws.cell(r, c).fill = FILL["S"]
        ws.row_dimensions[r].height = est_height([x.get("basis", "")], [58], 26, 90)
        r += 1
    ws.freeze_panes = "A4"


def sheet_verified(wb, data):
    rows = data.get("verified", [])
    if not rows:
        return
    ws = wb.create_sheet("검증완료")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "검증 완료 — 이상 없음"
    ws["A1"].font = Font(name=F, size=12, bold=True, color=NAVY)
    widths = [5, 26, 72, 40, 11]
    header(ws, 3, ["No", "검증 항목", "검증 내용", "대사 자료", "결과"], widths)
    r = 4
    for i, x in enumerate(rows, 1):
        put(ws, r, [x.get("no", i), x.get("item", ""), x.get("detail", ""),
                    x.get("source", ""), x.get("result", "일치")])
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 5).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 5).fill = FILL["V"]
        ws.row_dimensions[r].height = est_height([x.get("detail", "")], [72], 30, 90)
        r += 1
    ws.freeze_panes = "A4"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("findings")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()

    with open(a.findings, encoding="utf-8") as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    sheet_summary(wb, data)
    sheet_findings(wb, data)
    sheet_tieout(wb, data)
    sheet_verified(wb, data)
    wb.save(a.out)

    n = len(data.get("findings", []))
    print("지적 %d건 / 대사 %d행 / 검증완료 %d건  ->  %s"
          % (n, len(data.get("tieout", [])), len(data.get("verified", [])), a.out))


if __name__ == "__main__":
    main()
