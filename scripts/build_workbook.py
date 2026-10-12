# -*- coding: utf-8 -*-
"""검토·감사보고서용 검증 워크북 — 보고서를 엑셀에 옮기고 검증식을 살아 있는 수식으로 심는다.

사용법
    python build_workbook.py 01_report.json -o "풋팅결과_보고서명.xlsx"
        [--findings 03_findings.json] [--structure 04_structure.json] [--prev 01_prev.json]

입력
    01_report.json  extract.py 출력. "doc" (statements/notes) 가 있어야 한다 (dsd).
    03_findings.json  build_report.py 와 같은 스키마. 있으면 지적사항·검증완료 시트를 붙인다.
    04_structure.json 모델이 쓰는 구조 정의 (전부 선택).
        {"tol": 0.5,                          # 원 단위 허용오차
         "rules": {"PL": [{"target": "매출총이익", "expr": "매출액 - 매출원가"}]},
         "skip": {"PL": ["주당손익"]},          # 구조 미인식에서 제외
         "auto_rollforward": true,             # 주석 변동표(기초…기말) 행검증 자동
         "note_rules": [{"note": "6", "table": 3,
                         "cols": ["기초잔액","취득","처분","감가상각비"], "result": "기말잔액"}],
         "note_skip": [{"note": "18", "table": 4}],   # 표 단위로 푸팅검증 생략
         "xrefs": [{"label": "재고자산 평가액 합계 = BS 재고자산",
                    "from": {"note": "3", "table": 1, "row": "합계", "col": "평가액 당기"},
                    "to": {"sheet": "BS", "row": "재고자산", "period": 0},
                    "scale": 1000, "tol": 5000}],
         "wp": {"BS": {"file": "정산표.xlsx", "sheet": "BS", "label_col": "A", "first_row": 3,
                       "periods": [["B","C"],["D","E"]],
                       "sign_neg": ["대손충당금","감가상각누계액"],
                       "map": {"현금및현금성자산": ["제예금","기타제예금"]}}}}
    01_prev.json    전기 보고서의 extract 출력. 전기 열을 전기 보고서 당기 열과 대사한다.

원칙
    - 보고서에 적힌 합계는 값으로 그대로 둔다. 검증셀이 합계와 Σ항목을 비교한다.
      (합계 자리에 =SUM 을 넣으면 검증이 항상 TRUE 가 되어 의미가 없다.)
    - 검증셀 = IF(ABS(차이)<허용오차, TRUE, 차이). 숫자가 보이면 오류다.
    - 기계는 계층(ALEVEL)과 표준 산식만 안다. 그 밖의 소계는 '구조 미인식'으로 남겨 모델이 정한다.
"""
import argparse
import json
import os
import re
import sys

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

F = "맑은 고딕"
NAVY = "002060"
HDR = PatternFill("solid", fgColor="DDEBF7")
TOT = PatternFill("solid", fgColor="F2F2F2")
CHK = PatternFill("solid", fgColor="FFF2CC")
ERR = PatternFill("solid", fgColor="FFC7CE")
WPF = PatternFill("solid", fgColor="EDEDED")
_thin = Side(style="thin", color="BFBFBF")
BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
NUM = '#,##0;(#,##0);"-"'
STMT_NAME = {"BS": "재무상태표", "PL": "손익계산서", "MC": "제조원가명세서",
             "CE": "자본변동표", "CF": "현금흐름표", "RE": "결손금처리계산서"}


# ---------------------------------------------------------------- 라벨 정규화

_ROMAN = "ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩⅪⅫ"


def norm(s, keep_paren=False):
    """공백·번호·로마숫자·(주석n)·(손실) 등을 벗긴 비교용 라벨."""
    t = str(s or "")
    t = re.sub(r"\s+", "", t)
    t = re.sub(r"^[%sIVX]+\." % _ROMAN, "", t)
    t = re.sub(r"^[%s]+" % _ROMAN, "", t)
    t = re.sub(r"^\(?\d{1,2}[.)]", "", t)
    t = re.sub(r"^[가-힣]\.", "", t)
    t = re.sub(r"\(주석[\d,·\s]*\)", "", t)
    if not keep_paren:
        t = re.sub(r"\([^)]*\)", "", t)
    t = t.replace("및", "와").replace("의", "")
    t = t.replace("총손실", "총이익").replace("영업손실", "영업이익").replace("순손실", "순이익")
    t = t.replace("판매관리비", "판매비와관리비").replace("판매비와일반관리비", "판매비와관리비")
    return t.strip("*·.:- ")


def _safe(x, v):
    """문자열이 '='로 시작하면 수식이 아니라 텍스트로 저장한다 (보고서 문장 '= 1,726,196 ÷ 365' 등)."""
    x.value = v
    if isinstance(v, str) and v.startswith("="):
        x.data_type = "s"


def cell(ws, r, c, v=None, bold=False, fill=None, num=False, center=False, size=9, color=None, formula=False):
    x = ws.cell(r, c)
    if v is not None:
        if formula:
            x.value = v
        else:
            _safe(x, v)
    x.font = Font(name=F, size=size, bold=bold, color=color)
    x.border = BD
    if fill:
        x.fill = fill
    if num:
        x.number_format = NUM
        x.alignment = Alignment(horizontal="right", vertical="center")
    elif center:
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        x.alignment = Alignment(vertical="center")
    return x


def text(ws, r, c, v, size=9, bold=False, color=None):
    x = ws.cell(r, c)
    _safe(x, v)
    x.font = Font(name=F, size=size, bold=bold, color=color)
    x.alignment = Alignment(vertical="center")
    return x


def chk(ws, r, c, expr, tol):
    x = ws.cell(r, c, "=IF(ABS(%s)<%s,TRUE,%s)" % (expr, tol, expr))
    x.font = Font(name=F, size=9)
    x.fill = CHK
    x.border = BD
    x.number_format = NUM
    x.alignment = Alignment(horizontal="center", vertical="center")
    return x


class Ctx:
    """시트 하나의 검증셀 위치·집계."""

    def __init__(self, ws):
        self.ws = ws
        self.cells = {"foot": [], "rule": [], "xref": [], "wp": [], "prev": []}
        self.unrec = []      # 구조 미인식 라벨
        self.unmatched = []  # 조서 미매칭 라벨
        self.py = {"foot": [], "rule": [], "xref": [], "wp": [], "prev": []}  # 파이썬 재계산 차이
        self.fails = []  # (cat, desc, diff)
        self.next_row = 1

    def add(self, cat, r, c, diff=None, desc="", tol=0.5):
        self.cells[cat].append((r, c))
        if diff is not None:
            self.py[cat].append((diff, tol))
            if abs(diff) >= tol:
                self.fails.append((cat, "%s%d %s" % (CL(c), r, desc), diff))

    def ranges(self, cat):
        """같은 열에서 연속된 행을 묶어 A1 범위 목록으로."""
        pts = sorted(set(self.cells[cat]), key=lambda x: (x[1], x[0]))
        out = []
        for r, c in pts:
            if out and out[-1][1] == c and out[-1][3] == r - 1:
                out[-1][3] = r
            else:
                out.append([r, c, r, r])
        return ["%s%d:%s%d" % (CL(c), r0, CL(c), r1) for r0, c, _, r1 in out]

    def count_formula(self, cat):
        rs = self.ranges(cat)
        if not rs:
            return 0
        q = "'%s'!" % self.ws.title
        return "=" + "+".join("SUMPRODUCT(--ISNUMBER(%s%s))" % (q, r) for r in rs)

    def finish(self):
        for cat in self.cells:
            for rng in self.ranges(cat):
                first = rng.split(":")[0]
                self.ws.conditional_formatting.add(
                    rng, FormulaRule(formula=["ISNUMBER(%s)" % first], fill=ERR))


def _diff(a, b):
    if a is None or b is None:
        return None
    return round(a - b, 2)


# ---------------------------------------------------------------- 재무제표 계층·산식

def hierarchy(rows):
    """[(parent_idx, [child_idx...])] — ALEVEL 계층에서 소계와 직계 항목."""
    out = []
    n = len(rows)
    for i, r in enumerate(rows):
        L = r["level"]
        if L is None or not any(v is not None for v in r["vals"]):
            continue
        desc = []
        j = i + 1
        while j < n:
            lj = rows[j]["level"]
            if lj is None:
                j += 1
                continue
            if lj <= L:
                break
            desc.append(j)
            j += 1
        if not desc:
            continue
        m = min(rows[k]["level"] for k in desc)
        kids = [k for k in desc if rows[k]["level"] == m]
        if any(any(v is not None for v in rows[k]["vals"]) for k in kids):
            out.append((i, kids))
    return out


def find_row(rows, pat, start=0, end=None, need_val=None):
    end = len(rows) if end is None else end
    rx = re.compile(pat)
    for i in range(start, end):
        if rx.search(norm(rows[i]["label"])):
            if need_val is None or (any(v is not None for v in rows[i]["vals"]) == need_val):
                return i
    return None


def default_rules(key, rows):
    """[(target_idx, [(sign, idx)...], 설명)] — 표준 산식."""
    rules = []
    lvl0 = [i for i, r in enumerate(rows) if r["level"] == 0]

    def valrows(a, b):
        return [i for i in lvl0 if a < i < b and any(v is not None for v in rows[i]["vals"])]

    if key.startswith("BS"):
        h = {}
        for i, r in enumerate(rows):
            nl = norm(r["label"])
            if nl in ("자산", "부채", "자본") and not any(v is not None for v in r["vals"]):
                h[nl] = i
        ta = find_row(rows, r"^자산총계$")
        tl = find_row(rows, r"^부채총계$")
        te = find_row(rows, r"^자본총계$")
        tt = find_row(rows, r"^부채와자본총계$")
        if ta is not None and "자산" in h:
            rules.append((ta, [(1, i) for i in valrows(h["자산"], ta)], "자산총계 = Σ자산 대분류"))
        if tl is not None and "부채" in h:
            rules.append((tl, [(1, i) for i in valrows(h["부채"], tl)], "부채총계 = Σ부채 대분류"))
        if te is not None and "자본" in h:
            rules.append((te, [(1, i) for i in valrows(h["자본"], te)], "자본총계 = Σ자본 대분류"))
        if tt is not None and tl is not None and te is not None:
            rules.append((tt, [(1, tl), (1, te)], "부채와자본총계 = 부채총계 + 자본총계"))
    elif key.startswith("PL"):
        s = find_row(rows, r"^매출액$|^매출$|^수익$")
        c = find_row(rows, r"^매출원가$")
        g = find_row(rows, r"^매출총이익|^매출총손실")
        sga = find_row(rows, r"^판매비와관리비$|^판매비와일반관리비$")
        op = find_row(rows, r"^영업이익|^영업손실")
        oi = find_row(rows, r"^영업외수익$")
        oe = find_row(rows, r"^영업외비용$")
        pt = find_row(rows, r"^법인세비용차감전|^법인세차감전")
        tax = find_row(rows, r"^법인세비용$|^법인세비용\(수익\)$|^법인세등$")
        ni = find_row(rows, r"^(당기|반기|분기|중간)?순이익|^(당기|반기|분기)?순손실")
        if g is not None and s is not None and c is not None:
            rules.append((g, [(1, s), (-1, c)], "매출총이익 = 매출액 − 매출원가"))
        if op is not None and g is not None and sga is not None:
            rules.append((op, [(1, g), (-1, sga)], "영업이익 = 매출총이익 − 판매비와관리비"))
        if pt is not None and op is not None:
            t = [(1, op)]
            if oi is not None or oe is not None:
                if oi is not None:
                    t.append((1, oi))
                if oe is not None:
                    t.append((-1, oe))
                desc = "법인세비용차감전순이익 = 영업이익 + 영업외수익 − 영업외비용"
            else:
                # K-IFRS 단계식: 영업이익과 법인세차감전 사이의 최상위 항목(금융수익·금융비용·기타영업외수익·비용…)
                between = [i for i in lvl0 if op < i < pt and any(v is not None for v in rows[i]["vals"])]
                t += [((-1 if re.search(r"비용|손실|원가", norm(rows[i]["label"])) else 1), i) for i in between]
                desc = "법인세비용차감전순이익 = 영업이익 + 사이 수익 항목 − 사이 비용 항목"
            rules.append((pt, t, desc))
        if ni is not None and pt is not None:
            t = [(1, pt)]
            if tax is not None:
                t.append((-1, tax))
            rules.append((ni, t, "당기순이익 = 법인세비용차감전순이익 − 법인세비용"))
    elif key.startswith("MC") or key.startswith("CF"):
        pat = (r"^당기총제조비용$|^합계$|^당기제품제조원가$|^(당기|반기)제품제조원가$" if key.startswith("MC")
               else r"현금증가|현금감소|현금성자산증가|현금성자산감소|현금순증가|^기말현금")
        prev = None
        last = -1
        for i in lvl0:
            if re.search(pat, norm(rows[i]["label"])) and any(v is not None for v in rows[i]["vals"]):
                t = ([(1, prev)] if prev is not None else []) + [(1, k) for k in valrows(last, i)]
                if t:
                    rules.append((i, t, "%s = 직전 소계 + 사이 항목" % rows[i]["label"].strip()))
                prev, last = i, i
    return rules


def parse_expr(expr, rows):
    """'매출액 - 매출원가 + 기타' → [(sign, idx)]"""
    terms = []
    for m in re.finditer(r"([+-]?)\s*([^+\-]+)", expr):
        sign = -1 if m.group(1) == "-" else 1
        lab = m.group(2).strip()
        idx = None
        for i, r in enumerate(rows):
            if norm(r["label"]) == norm(lab):
                idx = i
                break
        if idx is None:
            raise SystemExit("산식 라벨을 찾지 못함: %r" % lab)
        terms.append((sign, idx))
    return terms


# ---------------------------------------------------------------- 조서 로딩

def load_wp(spec, base_dir):
    path = spec["file"]
    if not os.path.isabs(path):
        path = os.path.join(base_dir, path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[spec["sheet"]]
    lc = openpyxl.utils.column_index_from_string(spec.get("label_col", "A"))
    periods = [[openpyxl.utils.column_index_from_string(c) for c in p] for p in spec["periods"]]
    rows = []
    for r in range(spec.get("first_row", 1), ws.max_row + 1):
        lab = ws.cell(r, lc).value
        if lab is None or str(lab).strip() == "":
            continue
        vals = []
        for cols in periods:
            nums = [ws.cell(r, c).value for c in cols
                    if isinstance(ws.cell(r, c).value, (int, float)) and not isinstance(ws.cell(r, c).value, bool)]
            vals.append(float(nums[-1]) if nums else None)
        rows.append({"row": r, "label": str(lab), "norm": norm(lab, keep_paren=True),
                     "vals": vals, "used": False})
    return {"rows": rows, "sheet": spec["sheet"], "file": os.path.basename(path),
            "sign_neg": [norm(x) for x in spec.get("sign_neg", [])],
            "map": {norm(k): v for k, v in spec.get("map", {}).items()}}


ALIASES = {  # 보고서 표기 → 조서에서 흔히 쓰는 표기 (값 대사로 검증되므로 오매칭 위험은 낮다)
    "정부보조금": ["국고보조금"], "당기법인세자산": ["선납세금"], "평가손실충당금": ["재고자산평가충당금"],
    "산업재산권": ["특허권"], "급여": ["직원급여"], "임차료": ["지급임차료"], "세금과공과": ["세금과공과금"],
    "무형자산상각비": ["무형고정자산상각", "무형자산상각"],
    "이익잉여금": ["미처리결손금", "이월이익잉여금"], "매입채무": ["외상매입금", "지급어음"],
    "장기기타채권": ["회생정리채권", "장기미수금"], "보통주자본금": ["자본금"],
}


def _lab_eq(wn, nl):
    if wn == nl or wn.startswith(nl + "("):
        return True
    a, b = (wn, nl) if len(wn) >= len(nl) else (nl, wn)
    return len(b) >= 3 and a.startswith(b) and len(a) - len(b) <= 2


def wp_match(wp, label, cursor, _depth=0):
    """보고서 라벨 하나에 대응하는 조서 행(들). 순서 근접 우선."""
    nl = norm(label)
    if not nl:
        return []
    if nl in wp["map"]:
        found = []
        for sub in wp["map"][nl]:
            m = wp_match(wp, sub, cursor, _depth + 1)
            found.extend(m)
        return found
    cands = [w for w in wp["rows"] if not w["used"] and
             (_lab_eq(w["norm"], nl) or _lab_eq(norm(w["label"]), nl))]
    if not cands and _depth == 0 and nl in ALIASES:
        found = []
        for sub in ALIASES[nl]:
            found.extend(wp_match(wp, sub, cursor, _depth + 1))
        if found:
            return found
    if not cands:
        return []
    after = [w for w in cands if w["row"] >= cursor[0]]
    pick = (after or cands)[0]
    pick["used"] = True
    cursor[0] = pick["row"]
    return [pick]


# ---------------------------------------------------------------- 재무제표 시트

def write_statement(wb, st, structure, base_dir, prev_doc, tol):
    key = st["key"]
    ws = wb.create_sheet(key)
    ws.sheet_view.showGridLines = False
    ctx = Ctx(ws)
    rows = st["rows"]
    np_ = len(st["periods"])
    is_ce = key.startswith("CE")

    text(ws, 1, 1, st["title"], size=12, bold=True, color=NAVY)
    text(ws, 2, 1, " / ".join(st["periods"]) if not is_ce else "")
    text(ws, 3, 1, "(단위: %s)" % st["unit"] if st["unit"] else "")
    back_col = 2 + np_ + 8
    ws.cell(1, back_col, '=HYPERLINK("#총괄표!A1","◀ 총괄표")').font = Font(name=F, size=9, color="0563C1")

    # 조서·전기 매핑 준비
    wp = None
    wspec = (structure.get("wp") or {}).get(key)
    if wspec:
        wp = load_wp(wspec, base_dir)
    prev_st = None
    if prev_doc and np_ >= 2 and not is_ce:
        for p in prev_doc.get("statements", []):
            if p["key"] == key:
                prev_st = p
    # 열 배치
    c_lab = 1
    c_val = [2 + i for i in range(np_)]
    c = 2 + np_
    c_rowchk = None
    if is_ce:
        c_rowchk = c
        c += 1
    c += 1  # 빈 열
    c_wp = {}
    if wp:
        for i in range(np_):
            c_wp[i] = (c, c + 1)
            c += 2
        c += 1
    c_prev = None
    if prev_st:
        c_prev = (c, c + 1)
        c += 3
    c_note = c

    # 헤더
    hr = 5
    cell(ws, hr, c_lab, "과            목", bold=True, fill=HDR, center=True, color=NAVY)
    for i, p in enumerate(st["periods"]):
        cell(ws, hr, c_val[i], p.replace("\n", " "), bold=True, fill=HDR, center=True, color=NAVY)
    if is_ce:
        cell(ws, hr, c_rowchk, "행합검증\n(총계−Σ)", bold=True, fill=HDR, center=True, color=NAVY)
    for i, (a, b) in c_wp.items():
        cell(ws, hr, a, "조서 %s" % (st["periods"][i].split("(")[0].strip()), bold=True, fill=HDR, center=True, color=NAVY)
        cell(ws, hr, b, "차이", bold=True, fill=HDR, center=True, color=NAVY)
    if c_prev:
        cell(ws, hr, c_prev[0], "전기보고서\n당기 열", bold=True, fill=HDR, center=True, color=NAVY)
        cell(ws, hr, c_prev[1], "차이", bold=True, fill=HDR, center=True, color=NAVY)
    cell(ws, hr, c_note, "비        고", bold=True, fill=HDR, center=True, color=NAVY)
    ws.column_dimensions[CL(c_lab)].width = 44
    for i in range(np_):
        ws.column_dimensions[CL(c_val[i])].width = 18
    for a, b in c_wp.values():
        ws.column_dimensions[CL(a)].width = 18
        ws.column_dimensions[CL(b)].width = 13
    if c_prev:
        ws.column_dimensions[CL(c_prev[0])].width = 18
        ws.column_dimensions[CL(c_prev[1])].width = 13
    if c_rowchk:
        ws.column_dimensions[CL(c_rowchk)].width = 14
    ws.column_dimensions[CL(c_note)].width = 40
    ws.freeze_panes = ws.cell(hr + 1, 2)

    # 본문
    r0 = hr + 1
    row_of = {}
    cursor = [0]
    prev_cursor = [0]
    prev_rows = None
    if prev_st:
        prev_rows = [{"row": i, "label": p["label"], "norm": norm(p["label"]), "vals": p["vals"], "used": False}
                     for i, p in enumerate(prev_st["rows"])]
    tot_pat = re.compile(r"총계$|^합계$|당기총제조비용|제품제조원가|기말의현금|기말현금")
    for i, r in enumerate(rows):
        rr = r0 + i
        row_of[i] = rr
        lv = r["level"] or 0
        lab = ("    " * lv) + r["label"].replace("\n", " ")
        is_tot = bool(tot_pat.search(norm(r["label"]))) or (is_ce and lv == 0)
        has_val = any(v is not None for v in r["vals"])
        cell(ws, rr, c_lab, lab, bold=(lv == 0 and has_val), fill=TOT if is_tot else None)
        for j, v in enumerate(r["vals"]):
            cell(ws, rr, c_val[j], v, num=True, fill=TOT if is_tot else None, bold=is_tot)
            if j < len(r.get("dup", [])) and r["dup"][j]:
                text(ws, rr, c_note, "세부·합계 열에 값이 둘 다 있음")
        # 조서 대사
        if wp and has_val:
            m = wp_match(wp, r["label"], cursor)
            if not m:
                for j in range(np_):
                    if r["vals"][j] is not None:
                        cell(ws, rr, c_wp[j][0], "미매칭", fill=WPF, center=True)
                ctx.unmatched.append(r["label"].strip())
            else:
                neg = any(norm(r["label"]).startswith(s) for s in wp["sign_neg"])
                for j in range(np_):
                    vs = [w["vals"][j] for w in m if j < len(w["vals"]) and w["vals"][j] is not None]
                    if r["vals"][j] is None and not vs:
                        continue
                    wv = sum(vs) if vs else None
                    if wv is not None and neg and wv > 0:
                        wv = -wv
                    cell(ws, rr, c_wp[j][0], wv, num=True, fill=WPF)
                    if wv is not None and r["vals"][j] is not None:
                        expr = "%s%d-%s%d" % (CL(c_val[j]), rr, CL(c_wp[j][0]), rr)
                        chk(ws, rr, c_wp[j][1], expr, tol)
                        ctx.add("wp", rr, c_wp[j][1], _diff(r["vals"][j], wv), r["label"].strip())
                src = ", ".join("%s!%s%d" % (wp["sheet"], wspec.get("label_col", "A"), w["row"]) for w in m)
                text(ws, rr, c_note, "조서 " + src)
        # 전기 보고서 대사
        if prev_rows and r["vals"][1] is not None:
            cands = [p for p in prev_rows if not p["used"] and p["norm"] == norm(r["label"])]
            after = [p for p in cands if p["row"] >= prev_cursor[0]]
            if cands:
                p = (after or cands)[0]
                p["used"] = True
                prev_cursor[0] = p["row"]
                pv = p["vals"][0] if p["vals"] else None
                cell(ws, rr, c_prev[0], pv, num=True, fill=WPF)
                if pv is not None:
                    expr = "%s%d-%s%d" % (CL(c_val[1]), rr, CL(c_prev[0]), rr)
                    chk(ws, rr, c_prev[1], expr, tol)
                    ctx.add("prev", rr, c_prev[1], _diff(r["vals"][1], pv), r["label"].strip())
            else:
                cell(ws, rr, c_prev[0], "미매칭", fill=WPF, center=True)
        # 자본변동표 행합
        if is_ce and has_val and np_ >= 2:
            last = np_ - 1
            if r["vals"][last] is not None:
                items = "+".join("%s%d" % (CL(c_val[j]), rr) for j in range(last))
                expr = "%s%d-(%s)" % (CL(c_val[last]), rr, items)
                chk(ws, rr, c_rowchk, expr, tol)
                ctx.add("foot", rr, c_rowchk,
                        _diff(r["vals"][last], sum(v or 0 for v in r["vals"][:last])), "행합 " + r["label"].strip())

    # ---- 검증 블록
    rr = r0 + len(rows) + 1
    text(ws, rr, 1, "■ 푸팅 검증 (소계 − Σ직계 항목)", bold=True, color=NAVY)
    rr += 1
    covered = set()
    if is_ce:
        # 잔액행 = 직전 잔액행 + 사이 변동행
        bal = [i for i, r in enumerate(rows) if r["level"] == 0 and any(v is not None for v in r["vals"])]
        for a, b in zip(bal, bal[1:]):
            mids = [k for k in range(a + 1, b) if rows[k]["level"] and any(v is not None for v in rows[k]["vals"])]
            # '총포괄손익 합계' 처럼 사이 변동행들의 소계인 행은 빼야 이중 계산이 안 된다 (Σ앞 변동행 = 그 행이면 소계)
            def _is_subtotal(k, prev):
                if not re.search(r"합계|소계", norm(rows[k]["label"])) or not prev:
                    return False
                return all(abs((rows[k]["vals"][j] or 0) - sum(rows[p]["vals"][j] or 0 for p in prev)) < tol
                           for j in range(np_))
            kept = []
            for k in mids:
                if not _is_subtotal(k, kept):
                    kept.append(k)
            mids = kept
            text(ws, rr, 1, "%s = %s + 변동 (행 %d~%d)" % (rows[b]["label"].strip(), rows[a]["label"].strip(),
                                                      row_of[a], row_of[b]))
            for j in range(np_):
                if rows[b]["vals"][j] is None:
                    continue
                terms = "+".join(["%s%d" % (CL(c_val[j]), row_of[a])] +
                                 ["%s%d" % (CL(c_val[j]), row_of[k]) for k in mids])
                expr = "%s%d-(%s)" % (CL(c_val[j]), row_of[b], terms)
                chk(ws, rr, c_val[j], expr, tol)
                py = (rows[a]["vals"][j] or 0) + sum(rows[k]["vals"][j] or 0 for k in mids)
                ctx.add("foot", rr, c_val[j], _diff(rows[b]["vals"][j], py), rows[b]["label"].strip())
            rr += 1
    else:
        for p, kids in hierarchy(rows):
            covered.add(p)
            # 차감 표시(소계는 음수, 항목은 양수) 감지: 전 기간에서 Σ항목 = −소계 이면 부호 반전
            sign = 1
            pv = [(rows[p]["vals"][j], sum(rows[k]["vals"][j] or 0 for k in kids))
                  for j in range(np_) if rows[p]["vals"][j] is not None]
            if pv and all(abs(a + b) < tol for a, b in pv) and any(abs(a - b) >= tol for a, b in pv):
                sign = -1
            text(ws, rr, 1, "%s  (행 %d = %sΣ행 %d~%d)%s" % (rows[p]["label"].strip(), row_of[p],
                                                          "−" if sign < 0 else "",
                                                          row_of[kids[0]], row_of[kids[-1]],
                                                          "  ※ 차감 표시: 항목 부호 반전" if sign < 0 else ""))
            for j in range(np_):
                if rows[p]["vals"][j] is None:
                    continue
                terms = "+".join("%s%d" % (CL(c_val[j]), row_of[k]) for k in kids)
                expr = "%s%d-(%s(%s))" % (CL(c_val[j]), row_of[p], "-" if sign < 0 else "", terms)
                chk(ws, rr, c_val[j], expr, tol)
                py = sign * sum(rows[k]["vals"][j] or 0 for k in kids)
                ctx.add("foot", rr, c_val[j], _diff(rows[p]["vals"][j], py), rows[p]["label"].strip())
            rr += 1
        rr += 1
        text(ws, rr, 1, "■ 산식 검증 (비단순합)", bold=True, color=NAVY)
        rr += 1
        rules = default_rules(key, rows)
        for spec in (structure.get("rules") or {}).get(key, []):
            ti = find_row(rows, "^" + re.escape(norm(spec["target"])) + "$")
            if ti is None:
                raise SystemExit("rules[%s] target 을 찾지 못함: %r" % (key, spec["target"]))
            terms = parse_expr(spec["expr"], rows)
            rules = [x for x in rules if x[0] != ti] + [(ti, terms, "%s = %s" % (spec["target"], spec["expr"]))]
        for ti, terms, desc in rules:
            if not terms:
                continue
            covered.add(ti)
            covered.update(k for _, k in terms)
            text(ws, rr, 1, "%s  (행 %d)" % (desc, row_of[ti]))
            for j in range(np_):
                if rows[ti]["vals"][j] is None:
                    continue
                s = "".join(("+" if sg > 0 else "-") + "%s%d" % (CL(c_val[j]), row_of[k]) for sg, k in terms)
                expr = "%s%d-(%s)" % (CL(c_val[j]), row_of[ti], s)
                chk(ws, rr, c_val[j], expr, tol)
                py = sum(sg * (rows[k]["vals"][j] or 0) for sg, k in terms)
                ctx.add("rule", rr, c_val[j], _diff(rows[ti]["vals"][j], py), desc)
            rr += 1
        skip = [norm(x) for x in (structure.get("skip") or {}).get(key, [])]
        for i, r in enumerate(rows):
            if r["level"] == 0 and i not in covered and any(v is not None for v in r["vals"]):
                if norm(r["label"]) in skip:
                    continue
                ctx.unrec.append(r["label"].strip())
        if ctx.unrec:
            rr += 1
            text(ws, rr, 1, "■ 구조 미인식 — 산식이 정의되지 않은 소계 (04_structure.json rules 로 정의)", bold=True, color="C00000")
            rr += 1
            for lab in ctx.unrec:
                text(ws, rr, 1, "   " + lab)
                rr += 1
    ctx.next_row = rr + 1
    return ctx, {"row_of": row_of, "c_val": c_val, "rows": rows, "periods": st["periods"]}


# ---------------------------------------------------------------- 주석 시트

TOTAL_RX = re.compile(r"^(합계|소계|계|총계|차감계|순변동|순액|합계액|차감후금액|장부금액|장부가액|순장부금액|"
                      r"기말(잔액|금액|장부금액|장부가액)?|(당|전|반|분)?기말미처리결손금|미처리결손금|"
                      r"(당|전)기말(잔액|금액)?)$")


def col_header(tbl, ci):
    parts = []
    for h in tbl["header"]:
        if ci < len(h) and h[ci] and (not parts or parts[-1] != h[ci]):
            parts.append(h[ci])
    return " ".join(parts)


def find_col(tbl, spec):
    if isinstance(spec, int):
        return spec
    toks = [t for t in re.split(r"\s+", spec.strip()) if t]
    ncol = max(len(r) for r in tbl["rows"]) if tbl["rows"] else 0
    for ci in range(ncol):
        h = re.sub(r"\s+", "", col_header(tbl, ci))
        if all(re.sub(r"\s+", "", t) in h for t in toks):
            return ci
    return None


def find_note_row(tbl, spec):
    if isinstance(spec, int):
        return spec
    for ri, row in enumerate(tbl["rows"]):
        if row and norm(row[0]["text"]) == norm(spec):
            return ri
    return None


NONADD_RX = re.compile(r"율|비율|%|일수|일자|단가|주당|배수|기간|만기|이자율")
DECR_RX = re.compile(r"처분|감소|회수|상환|지급|환입|상각|매각|제거|차감|감자|소각|배당|사용|대체|손상|폐기|평가손실")


_GTOT_RX = re.compile(r"(합계|소계|총계)$|^계$")


def _row_total_label(row):
    """합계 판단에 쓸 라벨 — 첫 칸이 rowspan 복제면 그 행의 첫 '자기 텍스트' 칸."""
    for c in row:
        if not c.get("span") and c["text"].strip():
            return norm(c["text"])
    return ""


def _group_candidate(tbl, tri):
    """그룹합 해석 — 라벨만 있는 행('유동', '기타채권')이나 rowspan 묶음이 그룹을 열고 합계·소계 행이 닫는다.

    합계 = 같은 그룹 안 항목 Σ + 바로 아래 그룹의 합계행 Σ.  ('유동 합계' = 매출채권 + 기타채권 소계,
    '합 계' = 유동 합계 + 비유동 합계.)  해당 없으면 None.
    """
    rows = tbl["rows"]
    stack = [{"items": [], "subs": [], "rowspan": False}]          # [0] = 최상위
    for k, row in enumerate(rows[:tri + 1]):
        if not row:
            continue
        has_num = any(c["num"] is not None for c in row[1:])
        lab = _row_total_label(row)
        is_tot = has_num and bool(_GTOT_RX.search(lab))
        if k == tri:
            if not is_tot:
                return None
            g = stack[-1]
            items = sorted(g["items"] + g["subs"])
            return ("그룹합", items, None) if items else None
        next_span = k + 1 < len(rows) and bool(rows[k + 1] and rows[k + 1][0].get("span"))
        if is_tot:
            if len(stack) > 1:
                stack.pop()
                stack[-1]["subs"].append(k)
            else:
                stack[-1]["items"], stack[-1]["subs"] = [], []    # 최상위 합계가 중간에 오면 새 구간
        elif has_num:
            if not row[0].get("span") and next_span:                 # rowspan 묶음의 첫 행
                stack.append({"items": [k], "subs": [], "rowspan": True})
            else:
                stack[-1]["items"].append(k)
        elif lab:                                                     # 라벨만 있는 행 = 그룹 머리
            stack.append({"items": [], "subs": [], "rowspan": False})
        if stack[-1]["rowspan"] and not next_span and len(stack) > 1:  # 합계 없이 끝난 rowspan 묶음은 상위에 합친다
            g = stack.pop()
            stack[-1]["items"] += g["items"]
            stack[-1]["subs"] += g["subs"]
    return None


def _total_candidates(tbl, tri, totals_before, start):
    """합계 행 tri 에 대한 해석 후보. (설명, items, extra)  extra: [prev_idx] 또는 [("-", idx)]"""
    rows = tbl["rows"]
    items = [k for k in range(start, tri) if any(c["num"] is not None for c in rows[k][1:])]
    lab = norm(rows[tri][0]["text"])
    prev = totals_before[-1] if totals_before else None
    cands = []
    if items:
        cands.append(("Σ(행)", items, None))
    if prev is not None and items:
        cands.append(("직전 소계 + Σ(행)", items, [prev]))
    subs = [t for t in totals_before if norm(rows[t][0]["text"]) in ("소계", "계")]
    if lab in ("합계", "총계", "합계액") and len(subs) >= 2:
        cands.append(("Σ소계", subs, None))
    if lab == "차감계" and len(subs) >= 2:
        cands.append(("소계 − 소계", [subs[-2]], [("-", subs[-1])]))
    gc = _group_candidate(tbl, tri)
    if gc and all(gc[1] != c[1] for c in cands):
        cands.append(gc)
    return cands


def _eval_cand(tbl, tri, cand, ci):
    rows = tbl["rows"]
    _, items, extra = cand

    def v(k):
        return (rows[k][ci]["num"] or 0) if ci < len(rows[k]) else 0
    total = sum(v(k) for k in items)
    for e in (extra or []):
        if isinstance(e, tuple):
            total -= v(e[1])
        else:
            total += v(e)
    return total


def _cand_expr(tbl, tri, cand, ci, r0):
    _, items, extra = cand
    col = CL(ci + 1)
    if items == list(range(items[0], items[-1] + 1)):
        s = "SUM(%s%d:%s%d)" % (col, r0 + items[0], col, r0 + items[-1])
    else:
        s = "+".join("%s%d" % (col, r0 + k) for k in items)
    for e in (extra or []):
        if isinstance(e, tuple):
            s = "%s-%s%d" % (s, col, r0 + e[1])
        else:
            s = "%s%d+%s" % (col, r0 + e, s)
    return "%s%d-(%s)" % (col, r0 + tri, s)


def write_note(wb, note, structure, tol):
    name = note["no"]
    tol_r = structure.get("tol_rounded", 1.5)
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ctx = Ctx(ws)
    ws.column_dimensions["A"].width = 46
    for c in range(2, 12):
        ws.column_dimensions[CL(c)].width = 17
    text(ws, 1, 1, "%s. %s" % (note["no"], note["title"]), size=12, bold=True, color=NAVY)
    ws.cell(1, 10, '=HYPERLINK("#총괄표!A1","◀ 총괄표")').font = Font(name=F, size=9, color="0563C1")
    r = 3
    tmaps = {}  # table index(1-based) -> {"r0":, "cols":, "tbl":}
    ti = 0
    note_rules = [x for x in structure.get("note_rules", []) if str(x.get("note")) == str(note["no"])]
    note_skip = {(str(x.get("note")), int(x.get("table"))) for x in structure.get("note_skip", [])}
    auto_rf = structure.get("auto_rollforward", True)
    for it in note["items"]:
        if it["kind"] == "para":
            text(ws, r, 1, it["text"])
            r += 1
            continue
        ti += 1
        tbl = it
        t_tol = tol if (not tbl.get("unit") or re.fullmatch(r"\s*원\s*", tbl["unit"])) else tol_r
        if tbl.get("unit"):
            text(ws, r, 2, "(단위: %s)" % tbl["unit"])
            r += 1
        ncol = max([len(h) for h in tbl["header"]] + [len(x) for x in tbl["rows"]] + [1])
        for h in tbl["header"]:
            for ci in range(ncol):
                v = h[ci] if ci < len(h) else ""
                cell(ws, r, ci + 1, v.replace("\n", " "), bold=True, fill=HDR, center=True, color=NAVY)
            r += 1
        r0 = r
        totals = []
        for ri, row in enumerate(tbl["rows"]):
            lab = row[0]["text"] if row else ""
            is_tot = bool(TOTAL_RX.match(norm(lab)))
            for ci in range(ncol):
                c_ = row[ci] if ci < len(row) else {"text": "", "num": None}
                tx = c_["text"].strip()
                if ci > 0 and c_["num"] is not None and "%" not in tx:
                    cell(ws, r, ci + 1, c_["num"], num=True, fill=TOT if is_tot else None, bold=is_tot)
                elif ci > 0 and tx in ("-", "–", "—"):
                    # '-' 는 0 으로 저장 (표시형식이 0 을 '-' 로 보여 준다). 텍스트면 행검증 수식이 #VALUE! 가 된다.
                    cell(ws, r, ci + 1, 0, num=True, fill=TOT if is_tot else None, bold=is_tot)
                else:
                    cell(ws, r, ci + 1, tx.replace("\n", " "), fill=TOT if is_tot else None,
                         bold=is_tot, center=(ci > 0))
            if is_tot:
                totals.append(ri)
            r += 1
        tmaps[ti] = {"r0": r0, "tbl": tbl, "ncol": ncol}
        # 열 규칙 (변동표 행검증): 지정 or 자동
        rule = next((x for x in note_rules if int(x.get("table", 0)) == ti), None)
        rf_cols = None
        if rule and rule.get("cols"):
            rf_cols = ([find_col(tbl, c) for c in rule["cols"]], find_col(tbl, rule["result"]))
        elif auto_rf and tbl["header"]:
            heads = [re.sub(r"\s+", "", col_header(tbl, ci)) for ci in range(ncol)]
            b = next((i for i, h in enumerate(heads) if re.match(r"^(기초|당기초|전기초)", h)), None)
            e = next((i for i, h in enumerate(heads) if re.match(r"^(기말|당기말|전기말)", h) and i > (b or 0)), None)
            if b is not None and e is not None and e - b >= 2:
                rf_cols = (list(range(b, e)), e)
        if rf_cols and None not in rf_cols[0] and rf_cols[1] is not None:
            srcs, res = rf_cols

            # 부호 후보: ① 표시 부호 그대로 합산 ② 감소성 열(처분·회수·상환…)은 차감
            def _rowdiff(signs):
                out = []
                for row in tbl["rows"]:
                    rv = row[res]["num"] if res < len(row) else None
                    if rv is None:
                        continue
                    out.append(rv - sum(sg * ((row[s_]["num"] or 0) if s_ < len(row) else 0)
                                        for sg, s_ in zip(signs, srcs)))
                return out
            plain = [1] * len(srcs)
            signed = [(-1 if DECR_RX.search(re.sub(r"\s+", "", col_header(tbl, s_))) else 1) for s_ in srcs]
            signs = plain
            if signed != plain:
                bad_plain = sum(1 for d in _rowdiff(plain) if abs(d) >= t_tol)
                bad_signed = sum(1 for d in _rowdiff(signed) if abs(d) >= t_tol)
                if bad_signed < bad_plain:
                    signs = signed
            cc = ncol + 1
            desc = " ".join(("−" if sg < 0 else "+") + re.sub(r"\s+", "", col_header(tbl, s_))
                            for sg, s_ in zip(signs, srcs)).lstrip("+")
            cell(ws, r0 - 1, cc, "└ 행검증: %s − (%s)" % (re.sub(r"\s+", "", col_header(tbl, res)), desc),
                 bold=True, fill=HDR, center=True, color=NAVY)
            ws.column_dimensions[CL(cc)].width = 22
            for ri, row in enumerate(tbl["rows"]):
                rv = row[res]["num"] if res < len(row) else None
                if rv is None:
                    continue
                rr_ = r0 + ri
                terms = "".join(("-" if sg < 0 else "+") + "%s%d" % (CL(s_ + 1), rr_) for sg, s_ in zip(signs, srcs))
                expr = "%s%d-(%s)" % (CL(res + 1), rr_, terms.lstrip("+"))
                chk(ws, rr_, cc, expr, t_tol)
                py = sum(sg * ((row[s_]["num"] or 0) if s_ < len(row) else 0) for sg, s_ in zip(signs, srcs))
                ctx.add("foot", rr_, cc, _diff(rv, py), "표%d 행검증 %s" % (ti, row[0]["text"]), t_tol)
        # 합계 검증
        if (str(note["no"]), ti) not in note_skip:
            start = 0
            done = []
            for tri in totals:
                cands = _total_candidates(tbl, tri, done, start)
                if cands:
                    num_cols = [ci for ci in range(1, ncol)
                                if ci < len(tbl["rows"][tri]) and tbl["rows"][tri][ci]["num"] is not None
                                and not NONADD_RX.search(re.sub(r"\s+", "", col_header(tbl, ci)))]
                    # 전 열에서 오류가 가장 적은 해석을 고른다 (앞 후보 우선)
                    best, best_bad = cands[0], None
                    for cand in cands:
                        bad = sum(1 for ci in num_cols
                                  if abs(tbl["rows"][tri][ci]["num"] - _eval_cand(tbl, tri, cand, ci)) >= t_tol)
                        if best_bad is None or bad < best_bad:
                            best, best_bad = cand, bad
                        if bad == 0:
                            break
                    text(ws, r, 1, "  └ 푸팅검증: %s(행 %d) − %s"
                         % (tbl["rows"][tri][0]["text"], r0 + tri, best[0]))
                    for ci in num_cols:
                        tv = tbl["rows"][tri][ci]["num"]
                        expr = _cand_expr(tbl, tri, best, ci, r0)
                        chk(ws, r, ci + 1, expr, t_tol)
                        py = _eval_cand(tbl, tri, best, ci)
                        ctx.add("foot", r, ci + 1, _diff(tv, py),
                                "표%d %s [%s]" % (ti, tbl["rows"][tri][0]["text"], col_header(tbl, ci).replace(chr(10), " ")), t_tol)
                    r += 1
                done.append(tri)
                start = tri + 1
        r += 1
    ctx.next_row = r + 1
    return ctx, tmaps


# ---------------------------------------------------------------- 크로스 참조

def _stmt_cell(smaps, sheet, row_spec, period):
    m = smaps.get(sheet)
    if not m:
        return None, None
    rows = m["rows"]
    idx = None
    if isinstance(row_spec, int):
        idx = row_spec
    else:
        pat = "^" + re.escape(norm(row_spec)) + "$"
        idx = find_row(rows, pat, need_val=True)
        if idx is None:
            idx = find_row(rows, re.escape(norm(row_spec)), need_val=True)
    if idx is None or period >= len(m["c_val"]):
        return None, None
    return "'%s'!%s%d" % (sheet, CL(m["c_val"][period]), m["row_of"][idx]), rows[idx]["vals"][period]


def _note_cell(nmaps, note, table, row_spec, col_spec):
    tm = nmaps.get(str(note), {}).get(int(table))
    if not tm:
        return None, None
    tbl = tm["tbl"]
    ri = find_note_row(tbl, row_spec)
    ci = find_col(tbl, col_spec)
    if ri is None or ci is None:
        return None, None
    v = tbl["rows"][ri][ci]["num"] if ci < len(tbl["rows"][ri]) else None
    return "'%s'!%s%d" % (str(note), CL(ci + 1), tm["r0"] + ri), v


def resolve_ref(ref, smaps, nmaps):
    if "note" in ref:
        return _note_cell(nmaps, ref["note"], ref.get("table", 1), ref.get("row", "합계"), ref.get("col", 1))
    return _stmt_cell(smaps, ref["sheet"], ref["row"], ref.get("period", 0))


def builtin_xrefs(smaps):
    out = []
    if "CF" in smaps and "PL" in smaps:
        cf, pl = smaps["CF"]["rows"], smaps["PL"]["rows"]
        ni_cf = find_row(cf, r"순이익|순손실", need_val=True)
        ni_pl = find_row(pl, r"^(당기|반기|분기|중간)?순이익|^(당기|반기|분기)?순손실", need_val=True)
        if ni_cf is not None and ni_pl is not None:
            for p in range(min(len(smaps["CF"]["c_val"]), len(smaps["PL"]["c_val"]))):
                out.append({"label": "현금흐름표 순이익 = 손익계산서 순이익 (기간 %d)" % (p + 1),
                            "from": {"sheet": "CF", "row": ni_cf, "period": p},
                            "to": {"sheet": "PL", "row": ni_pl, "period": p}, "scale": 1})
    if "CF" in smaps and "BS" in smaps:
        cf, bs = smaps["CF"]["rows"], smaps["BS"]["rows"]
        ec = find_row(cf, r"^기말현금", need_val=True)
        cash = find_row(bs, r"^현금과현금성자산|^현금및현금성자산|^현금$", need_val=True)
        if ec is not None and cash is not None:
            out.append({"label": "현금흐름표 기말현금 = 재무상태표 현금및현금성자산 (당기)",
                        "from": {"sheet": "CF", "row": ec, "period": 0},
                        "to": {"sheet": "BS", "row": cash, "period": 0}, "scale": 1})
            if not re.search(r"반기|분기", " ".join(smaps["CF"]["periods"])) and len(smaps["BS"]["c_val"]) > 1:
                out.append({"label": "현금흐름표 기말현금 = 재무상태표 현금및현금성자산 (전기)",
                            "from": {"sheet": "CF", "row": ec, "period": 1},
                            "to": {"sheet": "BS", "row": cash, "period": 1}, "scale": 1})
    if "CE" in smaps and "BS" in smaps:
        ce, bs = smaps["CE"]["rows"], smaps["BS"]["rows"]
        tot_col = next((i for i, p in enumerate(smaps["CE"]["periods"]) if "총계" in re.sub(r"\s", "", p)), None)
        bal = [i for i, r in enumerate(ce) if r["level"] == 0 and any(v is not None for v in r["vals"])]
        te = find_row(bs, r"^자본총계$", need_val=True)
        if tot_col is not None and bal and te is not None:
            out.append({"label": "자본변동표 기말 총계 = 재무상태표 자본총계 (당기)",
                        "from": {"sheet": "CE", "row": bal[-1], "period": tot_col},
                        "to": {"sheet": "BS", "row": te, "period": 0}, "scale": 1})
        if "PL" in smaps:
            pl = smaps["PL"]["rows"]
            ni_pl = find_row(pl, r"^(당기|반기|분기|중간)?순이익|^(당기|반기|분기)?순손실", need_val=True)
            re_col = next((i for i, p in enumerate(smaps["CE"]["periods"])
                           if re.search(r"이익잉여금|결손금", re.sub(r"\s", "", p))), None)
            ni_rows = [i for i, r in enumerate(ce) if r["level"] and re.search(r"순이익|순손실", norm(r["label"]))]
            if ni_pl is not None and re_col is not None and ni_rows:
                for p, i in enumerate(reversed(ni_rows)):
                    if p >= len(smaps["PL"]["c_val"]):
                        break
                    out.append({"label": "자본변동표 순이익 = 손익계산서 순이익 (기간 %d)" % (p + 1),
                                "from": {"sheet": "CE", "row": i, "period": re_col},
                                "to": {"sheet": "PL", "row": ni_pl, "period": p}, "scale": 1})
    return out


def write_xrefs(wb, xrefs, smaps, nmaps, ctxs, tol):
    per_sheet = {}
    for x in xrefs:
        fr = x["from"]
        sheet = str(fr.get("note")) if "note" in fr else fr["sheet"]
        per_sheet.setdefault(sheet, []).append(x)
    for sheet, items in per_sheet.items():
        if sheet not in ctxs:
            print("  ! xref 대상 시트 없음:", sheet)
            continue
        ctx = ctxs[sheet]
        ws = ctx.ws
        r = ctx.next_row
        text(ws, r, 1, "■ 크로스 참조", bold=True, color=NAVY)
        r += 1
        for x in items:
            a, av = resolve_ref(x["from"], smaps, nmaps)
            b, bv = resolve_ref(x["to"], smaps, nmaps)
            label = x.get("label", "")
            if a is None or b is None:
                text(ws, r, 1, "%s — 참조를 찾지 못함 (%s)" % (label, "from" if a is None else "to"), color="C00000")
                r += 1
                continue
            scale = x.get("scale", 1)
            t = x.get("tol", tol if scale == 1 else scale * 5)
            a_local = a.replace("'%s'!" % sheet, "")
            expr = "%s*%s-%s" % (a_local, scale, b) if scale != 1 else "%s-%s" % (a_local, b)
            text(ws, r, 1, label)
            chk(ws, r, 2, expr, t)
            text(ws, r, 3, "→ %s" % b.replace("'", ""))
            py = None if av is None or bv is None else round(av * scale - bv, 2)
            ctx.add("xref", r, 2, py, label, t)
            r += 1
        ctx.next_row = r + 1


# ---------------------------------------------------------------- 서술 검산 (deal 보고서)

def write_claims(wb, claims, smaps, nmaps, tol):
    """보고서 문장 속 수치를 표 값에서 재계산한다.

    claims: [{"loc": "슬라이드 3", "text": "매출은 전년 대비 20% 증가", "claim": 20, "unit": "%",
              "expr": "({a}/{b}-1)*100", "refs": {"a": {"note":"S4","table":1,"row":"매출액","col":"2025"},
                                                  "b": {"note":"S4","table":1,"row":"매출액","col":"2024"}},
              "tol": 0.5}]
    refs 는 xrefs 의 from/to 와 같은 형식. expr 의 {키} 자리에 셀 주소가 들어간다.
    """
    ws = wb.create_sheet("서술검산")
    ws.sheet_view.showGridLines = False
    ctx = Ctx(ws)
    text(ws, 1, 1, "서술 검산 — 본문 문장의 수치를 표 값으로 재계산", size=12, bold=True, color=NAVY)
    ws.cell(1, 8, '=HYPERLINK("#총괄표!A1","◀ 총괄표")').font = Font(name=F, size=9, color="0563C1")
    heads = ["No", "위치", "문장", "기재값", "단위", "재계산식", "재계산값", "차이"]
    widths = [5, 12, 60, 13, 6, 46, 14, 12]
    for i, (h, w) in enumerate(zip(heads, widths)):
        cell(ws, 3, i + 1, h, bold=True, fill=HDR, center=True, color=NAVY)
        ws.column_dimensions[CL(i + 1)].width = w
    r = 4
    for i, c in enumerate(claims, 1):
        cell(ws, r, 1, i, center=True)
        cell(ws, r, 2, c.get("loc", ""))
        cell(ws, r, 3, c.get("text", ""))
        cell(ws, r, 4, c.get("claim"), num=isinstance(c.get("claim"), (int, float)))
        if isinstance(c.get("claim"), (int, float)) and c.get("unit") == "%":
            ws.cell(r, 4).number_format = "0.0"
        cell(ws, r, 5, c.get("unit", ""), center=True)
        expr = c.get("expr", "")
        pyvals = {}
        ok = True
        for k, ref in (c.get("refs") or {}).items():
            addr, val = resolve_ref(ref, smaps, nmaps)
            if addr is None:
                ok = False
                expr = expr.replace("{%s}" % k, "#참조없음")
                continue
            expr = expr.replace("{%s}" % k, addr)
            pyvals[k] = val
        cell(ws, r, 6, expr)
        if ok and expr:
            f = ws.cell(r, 7, "=" + expr)
            f.font = Font(name=F, size=9)
            f.border = BD
            f.number_format = "0.0" if c.get("unit") == "%" else NUM
            t = c.get("tol", 0.5)
            chk(ws, r, 8, "G%d-D%d" % (r, r), t)
            py = None
            try:
                e = c.get("expr", "")
                for k, v in pyvals.items():
                    e = e.replace("{%s}" % k, "(%r)" % (v if v is not None else float("nan")))
                py = eval(e, {"__builtins__": {}}, {})
                py = None if py != py else round(py - float(c.get("claim", 0)), 4)
            except Exception:
                py = None
            ctx.add("xref", r, 8, py, "서술 %s" % c.get("text", "")[:30], t)
        else:
            cell(ws, r, 7, "참조 없음", center=True, fill=ERR)
        r += 1
    ctx.next_row = r + 1
    return ctx


# ---------------------------------------------------------------- 총괄표

def write_summary(wb, meta, order, ctxs, findings, note_titles):
    ws = wb.create_sheet("총괄표", 0)
    ws.sheet_view.showGridLines = False
    text(ws, 2, 2, meta.get("title", "풋팅 검증 워크북"), size=14, bold=True, color=NAVY)
    text(ws, 3, 2, meta.get("subtitle", ""))
    text(ws, 4, 2, meta.get("note", ""))
    heads = ["시트", "제        목", "링크", "푸팅 오류", "산식 오류", "크로스 오류", "조서 차이",
             "전기 차이", "구조 미인식", "조서 미매칭"]
    widths = [8, 40, 14, 11, 11, 12, 11, 11, 12, 12]
    hr = 6
    for i, (h, w) in enumerate(zip(heads, widths)):
        cell(ws, hr, 2 + i, h, bold=True, fill=HDR, center=True, color=NAVY)
        ws.column_dimensions[CL(2 + i)].width = w
    r = hr + 1
    first = r
    for name in order:
        ctx = ctxs[name]
        title = STMT_NAME.get(name, note_titles.get(name, name))
        cell(ws, r, 2, name, center=True)
        cell(ws, r, 3, title)
        ws.cell(r, 4, '=HYPERLINK("#\'%s\'!A1","바로가기")' % name).font = Font(name=F, size=9, color="0563C1")
        ws.cell(r, 4).border = BD
        for j, cat in enumerate(["foot", "rule", "xref", "wp", "prev"]):
            cell(ws, r, 5 + j, ctx.count_formula(cat), num=True, formula=True)
        cell(ws, r, 10, len(ctx.unrec), num=True)
        cell(ws, r, 11, len(ctx.unmatched), num=True)
        r += 1
    cell(ws, r, 3, "합        계", bold=True, fill=TOT)
    for c in range(5, 12):
        cell(ws, r, c, "=SUM(%s%d:%s%d)" % (CL(c), first, CL(c), r - 1), num=True, bold=True, fill=TOT, formula=True)
    ws.conditional_formatting.add("%s%d:%s%d" % (CL(5), first, CL(9), r),
                                  FormulaRule(formula=["%s%d>0" % (CL(5), first)], fill=ERR))
    r += 2
    if findings:
        counts = {k: 0 for k in "SABC"}
        for f_ in findings.get("findings", []):
            counts[f_["sev"]] = counts.get(f_["sev"], 0) + 1
        text(ws, r, 2, "지적사항", bold=True, color=NAVY)
        r += 1
        for k, nm in (("S", "발행 전 필수 수정"), ("A", "내용 보강·확인"), ("B", "표 형식"), ("C", "문구·오타")):
            cell(ws, r, 2, k, center=True, bold=True)
            cell(ws, r, 3, nm)
            cell(ws, r, 4, counts[k], num=True)
            r += 1
        concl = findings.get("summary", {}).get("conclusion")
        if concl:
            r += 1
            text(ws, r, 2, "핵심 결론", bold=True, color=NAVY)
            r += 1
            c = ws.cell(r, 2, concl)
            c.font = Font(name=F, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
            ws.row_dimensions[r].height = min(300, 16 * (concl.count("\n") + 2))
            r += 1
    r += 1
    for l in ("※ 검증셀은 IF(ABS(차이)<허용오차, TRUE, 차이). 숫자가 보이면 오류이며 위 표는 그 개수를 센다.",
              "※ 보고서에 적힌 합계는 값으로 그대로 두고, 검증셀이 Σ항목과 비교한다.",
              "※ 크로스 참조 중 천원 표 ↔ 원 단위 본문은 단수차이를 고려해 ±5,000원을 기본 허용오차로 둔다.",
              "※ '구조 미인식'은 계층·표준 산식으로 검증하지 못한 소계 수다. 04_structure.json 의 rules 로 정의한다.",
              "※ '조서 차이'는 04_structure.json 의 wp 매핑이 있을 때만 채워진다.",
              "※ 천원 등 반올림 단위 표의 허용오차는 기본 ±1 (tol_rounded=1.5). 원 단위는 ±0 (tol=0.5).",
              "※ 주석 표 합계는 Σ행 · 직전소계+Σ행 · Σ소계 · 소계−소계 중 맞아떨어지는 해석을 골라 수식으로 남기고, 어느 해석인지 검증행에 적는다.",
              "※ 비율·일수·단가 열은 합산 검증에서 제외한다."):
        text(ws, r, 2, l, size=8)
        r += 1


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("report")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--findings")
    ap.add_argument("--structure")
    ap.add_argument("--prev")
    a = ap.parse_args()

    rep = json.load(open(a.report, encoding="utf-8"))
    doc = rep.get("doc")
    if not doc:
        raise SystemExit("01_report.json 에 doc 이 없다 (dsd·pptx 추출 결과만 지원).")
    structure = json.load(open(a.structure, encoding="utf-8")) if a.structure else {}
    findings = json.load(open(a.findings, encoding="utf-8")) if a.findings else None
    prev_doc = json.load(open(a.prev, encoding="utf-8")).get("doc") if a.prev else None
    tol = structure.get("tol", 0.5)
    base_dir = os.path.dirname(os.path.abspath(a.structure)) if a.structure else os.getcwd()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ctxs, smaps, nmaps, order = {}, {}, {}, []
    for st in doc["statements"]:
        ctx, m = write_statement(wb, st, structure, base_dir, prev_doc, tol)
        ctxs[st["key"]] = ctx
        smaps[st["key"]] = m
        order.append(st["key"])
    note_titles = {}
    for n in doc["notes"]:
        ctx, tm = write_note(wb, n, structure, tol)
        ctxs[n["no"]] = ctx
        nmaps[n["no"]] = tm
        note_titles[n["no"]] = "%s. %s" % (n["no"], n["title"])
        order.append(n["no"])

    xrefs = builtin_xrefs(smaps) + list(structure.get("xrefs", []))
    write_xrefs(wb, xrefs, smaps, nmaps, ctxs, tol)
    if structure.get("claims"):
        ctxs["서술검산"] = write_claims(wb, structure["claims"], smaps, nmaps, tol)
        note_titles["서술검산"] = "서술 검산 (본문 수치 재계산)"
        order.append("서술검산")
    for ctx in ctxs.values():
        ctx.finish()

    meta = {"title": (rep.get("meta", {}).get("company") or os.path.splitext(rep.get("name", ""))[0]) + " 풋팅 검증 워크북",
            "subtitle": os.path.basename(rep.get("source", "")),
            "note": "검증 대상: " + ", ".join(STMT_NAME.get(k, k) for k in smaps) +
                    " / 주석 %d개" % len(doc["notes"])}
    write_summary(wb, meta, order, ctxs, findings, note_titles)

    if findings:
        import build_report
        build_report.sheet_findings(wb, findings)
        build_report.sheet_tieout(wb, findings)
        build_report.sheet_verified(wb, findings)
        # 지적사항을 총괄표 바로 뒤로
        for nm in ("지적사항", "숫자대사", "검증완료"):
            if nm in wb.sheetnames:
                wb.move_sheet(nm, offset=-(len(wb.sheetnames) - 1 - wb.sheetnames.index(nm)) + 1)
    wb.save(a.out)

    # 콘솔 요약 (파이썬 재계산 기준)
    print("시트  푸팅  산식  크로스  조서  전기  | 미인식 미매칭   (오류 = 파이썬 재계산 차이 ≠ 0)")
    tot = {k: [0, 0] for k in ("foot", "rule", "xref", "wp", "prev")}
    for name in order:
        ctx = ctxs[name]
        parts = []
        for cat in ("foot", "rule", "xref", "wp", "prev"):
            n = len(ctx.cells[cat])
            bad = sum(1 for d, t_ in ctx.py[cat] if d is not None and abs(d) >= t_)
            tot[cat][0] += n
            tot[cat][1] += bad
            parts.append("%3d/%-3d" % (bad, n))
        print("%-4s %s | %d %d" % (name, " ".join(parts), len(ctx.unrec), len(ctx.unmatched)),
              ("  미인식: " + "; ".join(ctx.unrec)) if ctx.unrec else "")
    print("합계  " + " ".join("%3d/%-3d" % (b, n) for n, b in tot.values()))
    print()
    print("차이 상세 (시트 / 셀 항목 : 차이)")
    for name in order:
        for cat, desc, d in ctxs[name].fails:
            if cat == "prev" and a.prev == a.report:
                continue
            print("  [%s] %s %s : %s" % (name, cat, desc, "{:,.0f}".format(d)))
        if ctxs[name].unmatched:
            print("  [%s] 조서 미매칭: %s" % (name, "; ".join(ctxs[name].unmatched)))
    print("-> %s" % a.out)


if __name__ == "__main__":
    main()
